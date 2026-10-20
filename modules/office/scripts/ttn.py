"""Transport waybill (ТрН, ПП РФ №2200 прил. 4) from validated JSON to .xlsx; optional PDF via LibreOffice.

Deterministic layout only: values come from the input, empty fields stay empty, nothing is inferred.
This is the ПП 2200 «Транспортная накладная», not the 1-Т form Operator uses for event logistics.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
from collections import OrderedDict
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from jsonschema import Draft202012Validator, FormatChecker

VERSION = '1.0.0'
SCHEMA = Path(__file__).with_name('schemas') / 'ttn.schema.json'
FORM = {
    'form': 'Транспортная накладная, приложение N 4 к Правилам перевозок грузов автомобильным транспортом (ПП РФ от 21.12.2020 N 2200)',
    'edition': 'ред. ПП РФ от 30.11.2021 N 2116, действует с 01.03.2022',
    'rules_valid_until': '2026-09-01',
    'source': 'https://base.garant.ru/400111454/ (сверено 05.10.2026)',
    'note': 'С 01.09.2026 оформление перевозочных документов в электронном виде обязательно (Федеральный закон от 07.06.2025 N 140-ФЗ); бумажная ТрН — только в установленных исключениях. Перед выдачей сверить актуальную форму.',
}
OWNERSHIP = '1 - собственность; 2 - совместная собственность супругов; 3 - аренда; 4 - лизинг; 5 - безвозмездное пользование'
COLS = 8
WIDTH = 12.5


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            value.update(chunk)
    return value.hexdigest()


def unique_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError('Duplicate JSON key: ' + key)
        result[key] = value
    return result


def read_json(path: Path):
    def invalid_constant(value):
        raise ValueError('Non-finite JSON number: ' + value)
    return json.loads(path.read_text(encoding='utf-8-sig'), object_pairs_hook=unique_object, parse_constant=invalid_constant)


def schema_issues(data) -> list[dict]:
    schema = json.loads(SCHEMA.read_text(encoding='utf-8'))
    validator = Draft202012Validator(schema, format_checker=FormatChecker())
    return [{'code': 'schema', 'field': '/'.join(str(p) for p in error.absolute_path) or '$', 'detail': error.message}
            for error in sorted(validator.iter_errors(data), key=lambda e: list(map(str, e.absolute_path)))]


def dec(value) -> Decimal | None:
    return None if value in (None, '') else Decimal(str(value).replace(',', '.'))


def show_dec(value: Decimal | None) -> str:
    if value is None:
        return ''
    text = format(value.normalize(), 'f') if value == value.to_integral() else format(value, 'f')
    return text.replace('.', ',')


def show_date(value) -> str:
    return date.fromisoformat(value).strftime('%d.%m.%Y') if value else ''


def txt(*parts, sep='; ') -> str:
    return sep.join(str(p).strip() for p in parts if p is not None and p is not False and str(p).strip())


def check(data: dict) -> tuple[list[dict], list[dict]]:
    """Consistency checks on supplied values only; returns (issues, warnings)."""
    issues, warnings = [], []
    waybill_date = (data.get('waybill') or {}).get('date')
    if not waybill_date:
        warnings.append({'code': 'waybill_date_missing', 'field': 'waybill.date', 'detail': 'Дата накладной не указана'})
    elif date.fromisoformat(waybill_date) >= date.fromisoformat(FORM['rules_valid_until']):
        warnings.append({'code': 'form_edition_expired', 'field': 'waybill.date',
                         'detail': 'Дата не раньше 01.09.2026: Правила ПП 2200 действовали до 01.09.2026, обязательна ЭТрН; бумажная форма только для исключений — сверить актуальную форму'})
    required = {'shipper.details': (data.get('shipper') or {}).get('details'),
                'consignee.details': (data.get('consignee') or {}).get('details'),
                'cargo.lines': (data.get('cargo') or {}).get('lines'),
                'carrier.details': (data.get('carrier') or {}).get('details'),
                'carrier.drivers': (data.get('carrier') or {}).get('drivers'),
                'vehicle.reg_number': (data.get('vehicle') or {}).get('reg_number')}
    for field, value in required.items():
        if not value:
            warnings.append({'code': 'empty_field', 'field': field, 'detail': 'Обязательная по форме графа пуста'})
    vehicle = data.get('vehicle') or {}
    if vehicle.get('ownership_type') in (3, 4, 5) and not vehicle.get('ownership_docs'):
        warnings.append({'code': 'ownership_docs_missing', 'field': 'vehicle.ownership_docs',
                         'detail': 'Для типов владения 3, 4, 5 нужны реквизиты документа-основания'})
    cost = data.get('cost') or {}
    base, tax, total = dec(cost.get('without_tax')), dec(cost.get('tax_amount')), dec(cost.get('with_tax'))
    if None not in (base, tax, total) and base + tax != total:
        issues.append({'code': 'cost_mismatch', 'field': 'cost',
                       'detail': f'{show_dec(base)} + {show_dec(tax)} != {show_dec(total)}'})
    lines = (data.get('cargo') or {}).get('lines') or []
    places = [line.get('places') for line in lines]
    loading_places = (data.get('loading') or {}).get('places')
    if lines and None not in places and loading_places is not None and sum(places) != loading_places:
        issues.append({'code': 'places_mismatch', 'field': 'loading.places',
                       'detail': f'Сумма мест по строкам {sum(places)} != {loading_places}'})
    for i, line in enumerate(lines):
        gross, net = dec(line.get('gross_kg')), dec(line.get('net_kg'))
        if gross is not None and net is not None and net > gross:
            issues.append({'code': 'net_above_gross', 'field': f'cargo.lines.{i}', 'detail': 'Нетто больше брутто'})
    return issues, warnings


def cargo_groups(lines: list[dict]) -> list[dict]:
    """Blank rows: one per group (owner rule: groups in the blank, items in the inventory)."""
    groups: OrderedDict[str, list] = OrderedDict()
    for index, line in enumerate(lines, 1):
        groups.setdefault(line.get('group') or line['name'], []).append((index, line))
    result = []
    for name, items in groups.items():
        rows = [line for _, line in items]
        def total(key):
            values = [line.get(key) for line in rows]
            if None in values or '' in values:
                return None
            return sum(values) if key == 'places' else sum(dec(v) for v in values)
        numbers = [i for i, _ in items]
        ref = f'поз. {numbers[0]}' if len(numbers) == 1 else f'поз. {numbers[0]}–{numbers[-1]}' if numbers == list(range(numbers[0], numbers[-1] + 1)) else 'поз. ' + ', '.join(map(str, numbers))
        result.append({'name': name, 'ref': ref, 'count': len(rows), 'places': total('places'),
                       'gross': total('gross_kg'), 'net': total('net_kg'), 'volume': total('volume_m3'),
                       'value': total('declared_value'),
                       'state': txt(*OrderedDict.fromkeys(line.get('state_info') for line in rows)),
                       'marking': txt(*OrderedDict.fromkeys(line.get('marking') for line in rows)),
                       'packaging': txt(*OrderedDict.fromkeys(line.get('packaging') for line in rows)),
                       'extra': txt(*OrderedDict.fromkeys(txt(line.get('dimensions_m') and 'габариты ' + line['dimensions_m'],
                                                                line.get('density') and 'плотность ' + line['density'],
                                                                line.get('extra')) for line in rows)),
                       'dangerous': txt(*OrderedDict.fromkeys(line.get('dangerous_info') for line in rows))})
    return result


class Sheet:
    def __init__(self, ws, styles):
        self.ws, self.s, self.row = ws, styles, 1

    def height(self, text: str, chars: int, line_pt: float, minimum: float):
        lines = sum(max(1, -(-len(part) // chars)) for part in ('' if text is None else str(text)).split('\n'))
        return max(minimum, lines * line_pt)

    def put(self, col1, col2, value, style, height=None):
        ws = self.ws
        cell = ws.cell(self.row, col1, value if value not in ('', None) else None)
        for key, val in style.items():
            setattr(cell, key, val)
        if col2 > col1:
            for col in range(col1 + 1, col2 + 1):
                other = ws.cell(self.row, col)
                if 'border' in style:
                    other.border = style['border']
            ws.merge_cells(start_row=self.row, start_column=col1, end_row=self.row, end_column=col2)
        return cell

    def title(self, text):
        self.put(1, COLS, text, self.s['title'])
        self.ws.row_dimensions[self.row].height = 15
        self.row += 1

    def field(self, value, caption, col1=1, col2=COLS, advance=True):
        span = int((col2 - col1 + 1) * WIDTH * 1.05)
        self.put(col1, col2, value, self.s['value'])
        self.ws.row_dimensions[self.row].height = max(self.ws.row_dimensions[self.row].height or 0,
                                                       self.height(value, span, 11.5, 14))
        self.row += 1
        self.put(col1, col2, caption, self.s['caption'])
        self.ws.row_dimensions[self.row].height = max(self.ws.row_dimensions[self.row].height or 0,
                                                       self.height(caption, int(span * 1.5), 8.5, 9))
        if advance:
            self.row += 1
        else:
            self.row -= 1

    def pair(self, left, right):
        """Two side-by-side (value, caption) fields as in sections 1/1а and 12."""
        self.field(left[0], left[1], 1, COLS // 2, advance=False)
        self.field(right[0], right[1], COLS // 2 + 1, COLS)


def build_workbook(data: dict, issues: list, warnings: list, input_sha: str, input_path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style='thin', color='000000')
    styles = {
        'title': {'font': Font(name='Arial', size=9, bold=True), 'fill': PatternFill('solid', fgColor='E7E6E6'),
                  'alignment': Alignment(vertical='center', wrap_text=True), 'border': Border(top=thin, bottom=thin)},
        'value': {'font': Font(name='Arial', size=9), 'alignment': Alignment(wrap_text=True, horizontal='left', vertical='bottom'),
                  'border': Border(bottom=thin)},
        'caption': {'font': Font(name='Arial', size=6, italic=True), 'alignment': Alignment(wrap_text=True, horizontal='center', vertical='top')},
        'head': {'font': Font(name='Arial', size=12, bold=True), 'alignment': Alignment(vertical='center')},
        'label': {'font': Font(name='Arial', size=9), 'alignment': Alignment(horizontal='right', vertical='center')},
        'wrap': {'font': Font(name='Arial', size=8), 'alignment': Alignment(wrap_text=True, vertical='center')},
        'box': {'font': Font(name='Arial', size=9, bold=True), 'alignment': Alignment(horizontal='center', vertical='center'),
                'border': Border(top=thin, bottom=thin, left=thin, right=thin)},
    }
    get = lambda *keys: _dig(data, keys)
    wb = Workbook()
    ws = wb.active
    ws.title = 'ТрН'
    for col in range(1, COLS + 1):
        ws.column_dimensions[get_column_letter(col)].width = WIDTH
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = 'Транспортная накладная — стр. &P из &N'
    sheet = Sheet(ws, styles)

    def header(label, d, n):
        sheet.put(1, 3, label, styles['head'])
        sheet.put(4, 4, 'Дата', styles['label'])
        sheet.put(5, 5, d, styles['box'])
        sheet.put(6, 6, 'N', styles['label'])
        sheet.put(7, 8, n, styles['box'])
        ws.row_dimensions[sheet.row].height = 18
        sheet.row += 1

    header('Транспортная накладная', show_date(get('waybill', 'date')), get('waybill', 'number'))
    header('Заказ (заявка)', show_date(get('order', 'date')), get('order', 'number'))
    sheet.put(5, 6, 'Экземпляр N', styles['label'])
    sheet.put(7, 8, get('copy_number'), styles['box'])
    sheet.row += 2

    forwarder = get('shipper', 'is_forwarder')
    sheet.put(1, 4, '1. Грузоотправитель', styles['title'])
    sheet.put(5, COLS, '1а Заказчик услуг по организации перевозки груза (при наличии)', styles['title'])
    ws.row_dimensions[sheet.row].height = 24
    sheet.row += 1
    sheet.put(1, 3, 'является экспедитором', styles['label'])
    sheet.put(4, 4, 'X' if forwarder is True else None, styles['box'])
    sheet.row += 1
    sheet.pair((get('shipper', 'details'), '(реквизиты, позволяющие идентифицировать Грузоотправителя)'),
               (get('customer', 'details'), '(реквизиты, позволяющие идентифицировать Заказчика услуг по организации перевозки груза)'))
    sheet.field(get('customer', 'settlement_basis'), 'реквизиты документа, определяющего основания осуществления расчетов по договору перевозки иным лицом, отличным от грузоотправителя (при наличии)', 5, COLS)
    sheet.field(get('customer', 'contract'), '(реквизиты договора на выполнение услуг по организации перевозки груза)', 5, COLS)

    sheet.title('2. Грузополучатель')
    sheet.field(get('consignee', 'details'), '(реквизиты, позволяющие идентифицировать Грузополучателя)')
    sheet.field(get('consignee', 'delivery_address'), '(адрес места доставки груза)')

    lines = get('cargo', 'lines') or []
    groups = cargo_groups(lines)
    sheet.title('3. Груз')
    sheet.field('\n'.join([txt(f"{g['name']} ({g['ref']})", g['state'], sep=', ') for g in groups]
                          + (['Поштучная опись с серийными номерами — лист «Опись груза».'] if lines else [])),
                '(отгрузочное наименование груза (для опасных грузов - в соответствии с ДОПОГ), его состояние и другая необходимая информация о грузе)')
    sheet.field('\n'.join(txt(g['name'] + ':', g['places'] is not None and f"{g['places']} мест", g['marking'], g['packaging'], sep=' ')
                          for g in groups if g['places'] is not None or g['marking'] or g['packaging']),
                '(количество грузовых мест, маркировка, вид тары и способ упаковки)')
    sheet.field('\n'.join(txt(g['name'] + ':', g['gross'] is not None and f"брутто {show_dec(g['gross'])} кг",
                              g['net'] is not None and f"нетто {show_dec(g['net'])} кг",
                              g['volume'] is not None and f"объем {show_dec(g['volume'])} м3", g['extra'], sep=' ')
                          for g in groups if g['gross'] is not None or g['net'] is not None or g['volume'] is not None or g['extra']),
                '(масса груза брутто в килограммах, масса груза нетто в килограммах (при возможности ее определения), размеры (высота, ширина, длина) в метрах (при перевозке крупногабаритного груза), объем груза в кубических метрах и плотность груза в соответствии с документацией на груз (при необходимости), дополнительные характеристики груза, учитывающие отраслевые особенности (при необходимости)')
    sheet.field('\n'.join(filter(None, [get('cargo', 'dangerous_info')] + [txt(g['name'] + ':', g['dangerous'], sep=' ') for g in groups if g['dangerous']])),
                '(в случае перевозки опасного груза - информация по каждому опасному веществу, материалу или изделию в соответствии с пунктом 5.4.1 ДОПОГ)')
    values = [g['value'] for g in groups]
    sheet.field(f'{show_dec(sum(values))} руб.' if groups and None not in values else '\n'.join(
                    f"{g['name']}: {show_dec(g['value'])} руб." for g in groups if g['value'] is not None),
                '(объявленная стоимость (ценность) груза (при необходимости)')

    sheet.title('4. Сопроводительные документы на груз (при наличии)')
    sheet.field('; '.join(get('documents', 'regulated') or []), '(перечень прилагаемых к транспортной накладной документов, предусмотренных ДОПОГ, санитарными, таможенными (при наличии), карантинными, иными правилами в соответствии с законодательством Российской Федерации, либо регистрационные номера указанных документов, если такие документы (сведения о таких документах) содержатся в государственных информационных системах)')
    sheet.field('; '.join(get('documents', 'certificates') or []), '(перечень прилагаемых к грузу сертификатов, паспортов качества, удостоверений и других документов, наличие которых установлено законодательством Российской Федерации, либо регистрационные номера указанных документов, если такие документы (сведения о таких документах) содержатся в государственных информационных системах)')
    sheet.field('; '.join(get('documents', 'shipment_docs') or []), '(реквизиты, позволяющие идентифицировать документ (-ы), подтверждающий (-ие) отгрузку товаров) (при наличии), реквизиты сопроводительной ведомости (при перевозке груженых контейнеров или порожних контейнеров)')

    sheet.title('5. Указания грузоотправителя по особым условиям перевозки')
    sheet.field(get('instructions', 'route_and_terms'), '(маршрут перевозки, дата и время/сроки доставки груза (при необходимости)')
    sheet.field(get('instructions', 'readdress_contacts'), '(контактная информация о лицах, по указанию которых может осуществляться переадресовка)')
    sheet.field(get('instructions', 'regulatory'), '(указания, необходимые для выполнения фитосанитарных, санитарных, карантинных, таможенных и прочих требований, установленных законодательством Российской Федерации)')
    sheet.field(get('instructions', 'temperature_seals'), '(температурный режим перевозки груза (при необходимости), сведения о запорно-пломбировочных устройствах (в случае их предоставления грузоотправителем), запрещение перегрузки груза)')

    sheet.title('6. Перевозчик')
    sheet.field(get('carrier', 'details'), '(реквизиты, позволяющие идентифицировать Перевозчика)')
    sheet.field('\n'.join(txt(d['name'], d.get('license') and 'ВУ ' + d['license'], d.get('phone') and 'тел. ' + d['phone'], d.get('details'), sep=', ')
                          for d in get('carrier', 'drivers') or []),
                '(реквизиты, позволяющие идентифицировать водителя (-ей)')

    sheet.title('7. Транспортное средство')
    capacity, volume = dec(get('vehicle', 'capacity_t')), dec(get('vehicle', 'volume_m3'))
    sheet.field(txt(get('vehicle', 'type'), get('vehicle', 'brand'), capacity is not None and f'{show_dec(capacity)} т',
                    volume is not None and f'{show_dec(volume)} м3', sep=', '),
                '(тип, марка, грузоподъемность (в тоннах), вместимость (в кубических метрах)')
    sheet.field(txt(get('vehicle', 'reg_number'), get('vehicle', 'trailer') and 'прицеп (полуприцеп) ' + get('vehicle', 'trailer')),
                '(регистрационный номер транспортного средства)')
    sheet.put(1, 7, 'Тип владения: ' + OWNERSHIP, styles['wrap'])
    sheet.put(8, 8, get('vehicle', 'ownership_type'), styles['box'])
    ws.row_dimensions[sheet.row].height = 24
    sheet.row += 1
    sheet.field(get('vehicle', 'ownership_docs'), '(реквизиты документа (-ов), подтверждающего (- их) основание владения грузовым автомобилем (тягачом а также прицепом (полуприцепом) (для типов владения 3, 4, 5)')
    sheet.field(get('vehicle', 'special_permit'), '(номер, дата и срок действия специального разрешения, установленный маршрут движения тяжеловесного и (или) крупногабаритного транспортного средства или транспортного средства, перевозящего опасный груз) (при наличии)')

    sheet.title('8. Прием груза')
    sheet.field(get('loading', 'loader'), '(реквизиты лица, осуществляющего погрузку груза в транспортное средство)')
    sheet.field(get('loading', 'infrastructure_owner'), '(наименование (ИНН) владельца объекта инфраструктуры пункта погрузки)')
    sheet.field(get('loading', 'address'), '(адрес места погрузки)')
    sheet.field(get('loading', 'planned_at'), '(заявленные дата и время подачи транспортного средства под погрузку)')
    sheet.pair((get('loading', 'arrived_at'), '(фактические дата и время прибытия под погрузку)'),
               (get('loading', 'departed_at'), '(фактические дата и время убытия)'))
    gross = dec(get('loading', 'gross_kg'))
    sheet.field(txt(gross is not None and f'{show_dec(gross)} кг', get('loading', 'gross_method'), sep=', '),
                '(масса груза брутто в килограммах и метод ее определения (определение разницы между массой транспортного средства после погрузки и перед погрузкой по общей массе или взвешиванием поосно или расчетная масса груза)')
    sheet.pair((get('loading', 'places'), '(количество грузовых мест)'), (get('loading', 'packaging'), '(тара, упаковка (при наличии)'))
    sheet.field(get('loading', 'carrier_remarks'), '(оговорки и замечания перевозчика (при наличии) о дате и времени прибытия/убытия, о состоянии, креплении груза, тары, упаковки, маркировки, опломбирования, о массе груза и количестве грузовых мест, о проведении погрузочных работ)')
    sheet.pair((get('loading', 'loader_signatory'), '(подпись, расшифровка подписи лица, осуществившего погрузку груза, с указанием реквизитов документа, подтверждающего полномочия лица на погрузку груза)'),
               (get('loading', 'driver_signatory'), '(подпись, расшифровка подписи водителя, принявшего груз для перевозки)'))

    sheet.title('9. Переадресовка (при наличии)')
    sheet.field(get('readdress', 'date_and_kind'), '(дата, вид переадресовки на бумажном носителе или в электронном виде (с указанием вида доставки документа)')
    sheet.field(get('readdress', 'new_address_and_time'), '(адрес нового пункта выгрузки, новые дата и время подачи транспортного средства под выгрузку)')
    sheet.field(get('readdress', 'instructed_by'), '(реквизиты лица, от которого получено указание на переадресовку)')
    sheet.field(get('readdress', 'new_consignee'), '(при изменении получателя груза - реквизиты нового получателя)')

    sheet.title('10. Выдача груза')
    sheet.field(get('delivery', 'address'), '(адрес места выгрузки)')
    sheet.field(get('delivery', 'planned_at'), '(заявленные дата и время подачи транспортного средства под выгрузку)')
    sheet.pair((get('delivery', 'arrived_at'), '(фактические дата и время прибытия)'),
               (get('delivery', 'departed_at'), '(фактические дата и время убытия)'))
    sheet.field(get('delivery', 'condition'), '(фактическое состояние груза, тары, упаковки, маркировки, опломбирования)')
    dgross, dnet = dec(get('delivery', 'gross_kg')), dec(get('delivery', 'net_kg'))
    sheet.pair((get('delivery', 'places'), '(количество грузовых мест)'),
               (txt(dgross is not None and f'брутто {show_dec(dgross)} кг', dnet is not None and f'нетто {show_dec(dnet)} кг',
                    get('delivery', 'density') and 'плотность ' + get('delivery', 'density'), sep=', '),
                '(масса груза брутто в килограммах, масса груза нетто в килограммах (при возможности ее определения), плотность груза в соответствии с документацией на груз (при необходимости)'))
    sheet.field(get('delivery', 'carrier_remarks'), '(оговорки и замечания перевозчика (при наличии) о дате и времени прибытия/убытия, о состоянии груза, тары, упаковки, маркировки, опломбирования, о массе груза и количестве грузовых мест)')
    sheet.pair((get('delivery', 'consignee_signatory'), '(должность, подпись, расшифровка подписи грузополучателя или уполномоченного грузоотправителем лица)'),
               (get('delivery', 'driver_signatory'), '(подпись, расшифровка подписи водителя, сдавшего груз грузополучателю или уполномоченному грузополучателем лицу)'))

    sheet.title('11. Отметки грузоотправителей, грузополучателей, перевозчиков (при необходимости)')
    sheet.field(get('marks', 'description'), '(краткое описание обстоятельств, послуживших основанием для отметки, сведения о коммерческих и иных актах, в том числе о погрузке/выгрузке груза)')
    sheet.pair((get('marks', 'penalty'), '(расчет и размер штрафа)'), (get('marks', 'signature_date'), '(подпись, дата)'))

    sheet.title('12. Стоимость перевозки груза (установленная плата) в рублях (при необходимости)')
    sheet.pair((show_dec(dec(get('cost', 'without_tax'))), '(стоимость перевозки без налога - всего)'),
               (get('cost', 'tax_rate'), '(налоговая ставка)'))
    sheet.pair((show_dec(dec(get('cost', 'tax_amount'))), '(сумма налога, предъявляемая покупателю)'),
               (show_dec(dec(get('cost', 'with_tax'))), '(стоимость перевозки с налогом - всего)'))
    sheet.field(get('cost', 'calculation_method'), '(порядок (механизм) расчета (исчислений) платы) (при наличии порядка (механизма)')
    sheet.pair((get('cost', 'carrier', 'entity'), '(реквизиты, позволяющие идентифицировать Экономического субъекта, составляющего первичный учетный документ о факте хозяйственной жизни со стороны Перевозчика)'),
               (get('cost', 'shipper', 'entity'), '(реквизиты, позволяющие идентифицировать Экономического субъекта, составляющего первичный учетный документ о факте хозяйственной жизни со стороны Грузоотправителя)'))
    sheet.pair((get('cost', 'carrier', 'basis'), '(основание, по которому Экономический субъект является составителем документа о факте хозяйственной жизни)'),
               (get('cost', 'shipper', 'basis'), '(основание, по которому Экономический субъект является составителем документа о факте хозяйственной жизни)'))
    sheet.pair(('-', ''), (get('cost', 'payer'), '(реквизиты, позволяющие идентифицировать лицо, от которого будут поступать денежные средства)'))
    sheet.pair((get('cost', 'carrier', 'signatory'), '(подпись, расшифровка подписи лица, ответственного за оформление факта хозяйственной жизни со стороны Перевозчика (уполномоченного лица)'),
               (get('cost', 'shipper', 'signatory'), '(подпись, расшифровка подписи лица, ответственного за оформление факта хозяйственной жизни со стороны Грузоотправителя (уполномоченного лица)'))
    sheet.pair((get('cost', 'carrier', 'authority'), '(должность, основание полномочий физического лица, уполномоченного Перевозчиком (уполномоченным лицом), дата подписания)'),
               (get('cost', 'shipper', 'authority'), '(должность, основание полномочий физического лица, уполномоченного Грузоотправителем (уполномоченным лицом), дата подписания)'))
    ws.print_area = f'A1:{get_column_letter(COLS)}{sheet.row - 1}'

    inv = wb.create_sheet('Опись груза')
    heads = ['N', 'Группа', 'Наименование', 'Серийный N', 'Мест', 'Тара, упаковка', 'Маркировка',
             'Брутто, кг', 'Нетто, кг', 'Объем, м3', 'Объявл. стоимость, руб.', 'Примечание']
    widths = [5, 14, 40, 18, 7, 16, 14, 11, 11, 10, 14, 24]
    inv.append([f"Приложение к транспортной накладной N {get('waybill', 'number') or ''} от {show_date(get('waybill', 'date'))}".strip()])
    inv['A1'].font = Font(name='Arial', size=11, bold=True)
    inv.append(heads)
    for col, width in enumerate(widths, 1):
        inv.column_dimensions[get_column_letter(col)].width = width
        cell = inv.cell(2, col)
        cell.font = Font(name='Arial', size=9, bold=True)
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    numeric = {'gross_kg': 8, 'net_kg': 9, 'volume_m3': 10, 'declared_value': 11}
    for index, line in enumerate(lines, 1):
        row = [index, line.get('group'), line['name'], line.get('serial'), line.get('places'), line.get('packaging'),
               line.get('marking')] + [None] * 4 + [line.get('note')]
        inv.append(row)
        r = inv.max_row
        for key, col in numeric.items():
            value = dec(line.get(key))
            if value is not None:
                inv.cell(r, col, float(value)).number_format = '#,##0.###'
        for col in range(1, len(heads) + 1):
            cell = inv.cell(r, col)
            cell.font = Font(name='Arial', size=9)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    if lines:
        first, last = 3, inv.max_row
        inv.append(['', '', 'Итого', ''])
        r = inv.max_row
        inv.cell(r, 3).font = Font(name='Arial', size=9, bold=True)
        complete = {'places': 5, **numeric}
        for key, col in complete.items():
            if all(line.get(key) not in (None, '') for line in lines):
                letter = get_column_letter(col)
                cell = inv.cell(r, col, f'=SUM({letter}{first}:{letter}{last})')
                cell.font = Font(name='Arial', size=9, bold=True)
                cell.number_format = '#,##0.###'
    inv.freeze_panes = 'A3'
    inv.page_setup.paperSize = inv.PAPERSIZE_A4
    inv.page_setup.orientation = 'landscape'
    inv.sheet_properties.pageSetUpPr.fitToPage = True
    inv.page_setup.fitToWidth, inv.page_setup.fitToHeight = 1, 0
    inv.print_title_rows = '2:2'

    meta = wb.create_sheet('Служебное')
    meta.column_dimensions['A'].width, meta.column_dimensions['B'].width = 22, 110
    rows = [('Форма', FORM['form']), ('Редакция', FORM['edition']),
            ('Правила действуют до', show_date(FORM['rules_valid_until'])), ('Источник формы', FORM['source']),
            ('Примечание', FORM['note']), ('Генератор', f'office-pipeline ttn.py {VERSION}'),
            ('Вход', input_path), ('Вход SHA-256', input_sha),
            ('Создано (UTC)', datetime.now(timezone.utc).isoformat(timespec='seconds')),
            ('Статус', 'needs_review: черновик, не подписан; пустые графы заполняются вручную')]
    rows += [('Замечание', f"{w['code']}: {w['field']} — {w['detail']}") for w in warnings]
    rows += [('Ошибка', f"{i['code']}: {i['field']} — {i['detail']}") for i in issues]
    for key, value in rows:
        meta.append([key, value])
        meta.cell(meta.max_row, 2).alignment = Alignment(wrap_text=True, vertical='top')
    meta.sheet_state = 'hidden'
    return wb


def _dig(data, keys):
    for key in keys:
        if not isinstance(data, dict):
            return None
        data = data.get(key)
    return data


def soffice() -> str | None:
    found = shutil.which('soffice')
    for candidate in (found, r'C:\Program Files\LibreOffice\program\soffice.exe', r'C:\Program Files (x86)\LibreOffice\program\soffice.exe'):
        if candidate and Path(candidate).is_file():
            return candidate
    return None


def to_pdf(xlsx: Path, pdf: Path) -> None:
    binary = soffice()
    if binary is None:
        raise ValueError('LibreOffice (soffice) not found; PDF not created')
    with tempfile.TemporaryDirectory(prefix='ttn-pdf-') as tmp:
        profile = Path(tmp, 'profile').as_uri()
        run = subprocess.run([binary, f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'pdf',
                              '--outdir', tmp, str(xlsx)], capture_output=True, timeout=180)
        produced = Path(tmp, xlsx.stem + '.pdf')
        if run.returncode != 0 or not produced.is_file():
            raise ValueError(f'LibreOffice conversion failed (exit {run.returncode})')
        with produced.open('rb') as source, pdf.open('xb') as target:
            shutil.copyfileobj(source, target)


def main(argv=None):
    parser = argparse.ArgumentParser(description='Build a draft ТрН (ПП РФ №2200, прил. 4) .xlsx from validated JSON; never fills absent values.')
    parser.add_argument('input', type=Path, help='Input JSON (schemas/ttn.schema.json)')
    parser.add_argument('--output', type=Path, required=True, help='New .xlsx file; existing files are never overwritten.')
    parser.add_argument('--pdf', action='store_true', help='Also write <output>.pdf via LibreOffice (headless, isolated profile).')
    parser.add_argument('--check', action='store_true', help='Validate only; write nothing.')
    args = parser.parse_args(argv)
    try:
        source = args.input.resolve()
        output = args.output.resolve()
        if output.suffix.lower() != '.xlsx':
            raise ValueError('Output must be a new .xlsx file')
        pdf = output.with_suffix('.pdf')
        for path in (output, pdf) if args.pdf else (output,):
            if path.exists():
                raise ValueError('Output exists, refusing to overwrite: ' + str(path))
        data = read_json(source)
        found = schema_issues(data)
        if found:
            print(json.dumps({'input': str(source), 'validation': 'schema_failed', 'issues': found[:50]}, ensure_ascii=False))
            return 2
        issues, warnings = check(data)
        receipt = {'input': str(source), 'input_sha256': digest(source), 'validation': 'passed' if not issues else 'issues',
                   'quality': 'needs_review', 'form': 'ПП2200-прил4-ред2116', 'rules_valid_until': FORM['rules_valid_until'],
                   'issues': issues, 'warnings': warnings, 'tool_version': VERSION}
        if not args.check:
            workbook = build_workbook(data, issues, warnings, receipt['input_sha256'], str(source))
            output.parent.mkdir(parents=True, exist_ok=True)
            with output.open('xb') as stream:
                workbook.save(stream)
            receipt.update(output=str(output), sha256=digest(output))
            if args.pdf:
                to_pdf(output, pdf)
                receipt.update(pdf=str(pdf), pdf_sha256=digest(pdf))
        print(json.dumps(receipt, ensure_ascii=False))
        return 0 if not issues else 2
    except (OSError, UnicodeError, ValueError, subprocess.SubprocessError) as exc:
        print(json.dumps({'error': str(exc)}, ensure_ascii=False))
        return 1


if __name__ == '__main__':
    import sys
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8', errors='replace')
    raise SystemExit(main())
