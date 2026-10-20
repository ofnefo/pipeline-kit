"""Prepare a draft proposal and workbook inputs from an explicit project estimate.

Run with the optional office dependencies. This writes a draft DOCX and
workbook input JSON. Author/recalculate XLSX separately, then record QA.
"""
from __future__ import annotations
import argparse
import hashlib
import importlib.util
import json
import sys
from datetime import datetime, timezone
from pathlib import Path


def read(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def local(root, name):
    if not isinstance(name, str) or not name.strip() or Path(name).is_absolute():
        raise ValueError('Expected non-empty project-relative path')
    path = (root / name).resolve()
    if not path.is_relative_to(root.resolve()):
        raise ValueError('Path escapes project')
    return path


def money(value):
    return f'{float(value):,.2f}'.replace(',', ' ').replace('.', ',')


def prepare(project: Path, context_name: str, output_name: str):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    project = project.resolve()
    config = read(project / 'OFFICE_PROJECT.json')
    estimate_path = local(project, config['files']['estimate_json'])
    context_path = local(project, context_name)
    estimate, context = read(estimate_path), read(context_path)
    for field in ('title', 'intro', 'schedule', 'validity', 'tax_note'):
        if not isinstance(context.get(field), str) or not context[field].strip():
            raise ValueError(f'Proposal context requires {field}; do not invent missing facts')
    for field in ('assumptions', 'exclusions', 'sources'):
        if not isinstance(context.get(field), list) or not context[field] or any(not isinstance(x, str) or not x.strip() for x in context[field]):
            raise ValueError(f'Proposal context requires non-empty {field} strings')
    source_paths = [local(project, name) for name in context['sources']]
    if not all(path.is_file() for path in source_paths):
        raise ValueError('A proposal source is missing')
    spec = importlib.util.spec_from_file_location('proposal_oracle', Path(__file__).parent / 'estimate_engine.py')
    engine = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(engine)
    result = engine.json_ready(engine.calculate(estimate))
    if result['warnings'] or not result['items']:
        raise ValueError(f'Resolve estimate warnings before preparing this packet: {result["warnings"]}')
    if result['credit_accruals']:
        raise ValueError('This packet recipe does not author credit schedules; use the full Spreadsheets workflow')
    if any(row['amount_mode'] != 'formula' for row in result['items']):
        raise ValueError('This packet recipe requires formula rows; route explicit overrides through Spreadsheets')
    out = local(project, output_name)
    # One packet is one version; never merge a retry into an earlier packet.
    out.mkdir(parents=True, exist_ok=False)
    provenance = {path.relative_to(project).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
                  for path in [estimate_path, context_path, *source_paths]}
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        document.styles[name].font.name = 'Arial'
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for border in document.styles[name]._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    document.styles['Title'].font.size = Pt(23)
    document.styles['Heading 1'].font.size = Pt(13)
    document.add_paragraph(context['title'], 'Title')
    document.add_paragraph(f"{config['project']['name']} · {estimate['version']['id']} · Черновик", 'Subtitle')
    document.add_paragraph(context['intro'])
    document.add_heading('Состав работ', 1)
    table = document.add_table(rows=1, cols=3)
    table.style = 'Light Shading Accent 1'
    for cell, text in zip(table.rows[0].cells, ('Работы', 'Количество', 'Дни')):
        cell.text = text
    header = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for row in result['items']:
        cells = table.add_row().cells
        cells[0].text = row['name']
        cells[1].text = f"{row['quantity']} {row['unit']}"
        cells[2].text = row['days']
    for row in table.rows:
        for cell, width in zip(row.cells, (11, 3.5, 2.5)):
            cell.width = Cm(width)
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
    document.add_heading('Стоимость и сроки', 1)
    p = document.add_paragraph()
    p.add_run(f"Общая стоимость — {money(result['financial_result']['contract_price'])} {config['project']['currency']}.").bold = True
    document.add_paragraph(context['tax_note'])
    document.add_paragraph(context['schedule'])
    document.add_paragraph(context['validity'])
    for title, field in (('Условия расчёта', 'assumptions'), ('Исключения', 'exclusions')):
        document.add_heading(title, 1)
        for text in context[field]:
            document.add_paragraph(text, 'List Bullet')
    footer = section.footer.paragraphs[0]
    footer.text = 'Проект предложения для обсуждения. Требует согласования.'
    footer.runs[0].font.size = Pt(8)
    document.core_properties.author = ''
    document.core_properties.title = context['title']
    document.save(out / 'proposal.docx')
    (out / 'calculation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    (out / 'workbook-input.json').write_text(json.dumps({'project': config['project'], 'estimate': estimate, 'calculation': result,
                                                     'tax_note': context['tax_note'], 'sources': context['sources']}, ensure_ascii=False, indent=2), encoding='utf-8')
    (out / 'packet.json').write_text(json.dumps({'schema_version': 1, 'status': 'draft_needs_artifact_qa',
                                              'created_utc': datetime.now(timezone.utc).isoformat(), 'sources': provenance,
                                              'estimate_version': estimate['version']['id'], 'approval': 'missing'}, ensure_ascii=False, indent=2), encoding='utf-8')
    return out


def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--project', type=Path, required=True)
    parser.add_argument('--context', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    try:
        print(prepare(args.project, args.context, args.output))
        return 0
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 1

if __name__ == '__main__':
    raise SystemExit(main())
